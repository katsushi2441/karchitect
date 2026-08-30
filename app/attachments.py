"""参考資料として添付されたファイルから、設計に使えるテキストを取り出す。

なぜ「抽出して要約する」形にしたか（2026-08-30）:
  1ターンあたり約9,000トークン（SYSTEM 612 + スキーマ 2,587 + 要件JSON + 会話履歴）を
  すでに送っている。ここへ添付の全文を毎回足すと、A4十数ページのPDFやExcelの表で
  簡単に数万トークンへ膨らみ、応答が返らなくなる。
  そこで、アップロード時に一度だけ全文をAIに読ませて要点へ圧縮し、
  以降の会話では圧縮したものだけを添える。原文はサーバーに残し、必要なときだけ開く。

対応形式: PDF(テキスト埋め込みのみ) / Word(.docx) / Excel(.xlsx) / テキスト系
スキャン画像のPDFはOCRが必要なため、対応せずエラーとして返す。
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

# 上限。ここを超える資料は、設計の材料としてはそもそも大きすぎる。
MAX_FILE_BYTES = 10 * 1024 * 1024      # 1ファイル10MB
MAX_FILES_PER_PROJECT = 5              # 1プロジェクト5ファイル
MAX_EXCEL_CELLS = 10_000               # これを超えたら先頭だけ読む
MAX_EXTRACT_CHARS = 120_000            # 抽出テキストの上限（要約前）

ALLOWED_SUFFIXES = {".pdf", ".docx", ".xlsx", ".txt", ".md", ".csv"}


class AttachmentError(Exception):
    """利用者にそのまま見せてよいエラー。"""


@dataclass
class Extracted:
    text: str
    note: str = ""      # 打ち切りなど、利用者に伝えるべきこと


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t　]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _truncate(text: str) -> Extracted:
    if len(text) <= MAX_EXTRACT_CHARS:
        return Extracted(text=text)
    return Extracted(
        text=text[:MAX_EXTRACT_CHARS],
        note=f"文字数が多いため先頭{MAX_EXTRACT_CHARS:,}文字だけを読み込みました。",
    )


def _from_pdf(path: Path) -> Extracted:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise AttachmentError("PDFを読むための部品が入っていません。requirements.txt を入れ直してください。") from exc
    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        raise AttachmentError("PDFを開けませんでした。壊れているか、パスワードが掛かっている可能性があります。") from exc
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")
    text = _clean("\n\n".join(pages))
    if len(text) < 20:
        # 画像として取り込まれたPDF。OCRが要るので、ここでは受け付けない。
        raise AttachmentError(
            "このPDFからは文字を取り出せませんでした。"
            "紙をスキャンした画像のPDFには対応していません。"
            "元のWord/Excelファイル、またはテキストが選択できるPDFをお使いください。"
        )
    return _truncate(text)


def _from_docx(path: Path) -> Extracted:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise AttachmentError("Wordを読むための部品が入っていません。requirements.txt を入れ直してください。") from exc
    try:
        doc = Document(str(path))
    except Exception as exc:
        raise AttachmentError("Wordファイルを開けませんでした。.docx 形式かご確認ください（古い .doc には対応していません）。") from exc
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:          # 表は設計の材料になることが多いので拾う
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    text = _clean("\n".join(parts))
    if not text:
        raise AttachmentError("このWordファイルには読み取れる文字がありませんでした。")
    return _truncate(text)


def _from_xlsx(path: Path) -> Extracted:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise AttachmentError("Excelを読むための部品が入っていません。requirements.txt を入れ直してください。") from exc
    try:
        book = load_workbook(str(path), read_only=True, data_only=True)
    except Exception as exc:
        raise AttachmentError("Excelファイルを開けませんでした。.xlsx 形式かご確認ください（古い .xls には対応していません）。") from exc
    lines: list[str] = []
    cells = 0
    truncated = False
    has_value = False        # シート名だけで「中身あり」と誤判定しないための旗
    for sheet in book.worksheets:
        sheet_lines: list[str] = []
        for row in sheet.iter_rows(values_only=True):
            values = ["" if v is None else str(v).strip() for v in row]
            cells += len(values)
            if any(values):
                sheet_lines.append(" | ".join(values))
                has_value = True
            if cells >= MAX_EXCEL_CELLS:
                truncated = True
                break
        if sheet_lines:
            lines.append(f"## シート: {sheet.title}")
            lines.extend(sheet_lines)
        if truncated:
            break
    book.close()
    text = _clean("\n".join(lines))
    if not has_value or not text:
        raise AttachmentError("このExcelファイルには読み取れる値がありませんでした。")
    result = _truncate(text)
    if truncated:
        note = f"セル数が多いため先頭{MAX_EXCEL_CELLS:,}セルだけを読み込みました。"
        result.note = (result.note + " " + note).strip()
    return result


def _from_text(path: Path) -> Extracted:
    raw = path.read_bytes()
    for encoding in ("utf-8", "cp932", "euc-jp"):
        try:
            return _truncate(_clean(raw.decode(encoding)))
        except UnicodeDecodeError:
            continue
    raise AttachmentError("文字コードを判別できませんでした。UTF-8で保存し直してお試しください。")


_EXTRACTORS = {
    ".pdf": _from_pdf,
    ".docx": _from_docx,
    ".xlsx": _from_xlsx,
    ".txt": _from_text,
    ".md": _from_text,
    ".csv": _from_text,
}


def check_filename(filename: str) -> str:
    """拡張子を確かめ、小文字で返す。対応外はここで止める。"""
    suffix = Path(filename or "").suffix.lower()
    if suffix not in ALLOWED_SUFFIXES:
        raise AttachmentError(
            "対応していない形式です。PDF・Word(.docx)・Excel(.xlsx)・テキスト(.txt/.md/.csv)をお使いください。"
            "（古い .doc / .xls、スキャンした画像のPDFには対応していません）"
        )
    return suffix


def extract(path: Path, filename: str) -> Extracted:
    suffix = check_filename(filename)
    return _EXTRACTORS[suffix](path)
