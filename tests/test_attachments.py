"""参考資料（添付ファイル）の読み取りテスト。

実ファイルをその場で作って読ませる。ライブラリの挙動に依存する箇所なので、
モックではなく本物を通す。
"""
from __future__ import annotations

from pathlib import Path

import pytest

from app import attachments


def test_対応外の拡張子は受け付けない():
    with pytest.raises(attachments.AttachmentError) as exc:
        attachments.check_filename("設計書.doc")
    assert "対応していない" in str(exc.value)


def test_拡張子は小文字で返す():
    assert attachments.check_filename("資料.PDF") == ".pdf"


def test_テキストを読める(tmp_path: Path):
    f = tmp_path / "memo.txt"
    f.write_text("受注から請求までの流れ\n担当は経理", encoding="utf-8")
    got = attachments.extract(f, "memo.txt")
    assert "受注から請求" in got.text


def test_cp932のテキストも読める(tmp_path: Path):
    f = tmp_path / "memo.txt"
    f.write_bytes("見積書の項目".encode("cp932"))
    assert "見積書" in attachments.extract(f, "memo.txt").text


def test_wordの本文と表を読める(tmp_path: Path):
    from docx import Document

    doc = Document()
    doc.add_paragraph("業務フローの概要")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "項目"
    table.cell(0, 1).text = "内容"
    table.cell(1, 0).text = "顧客名"
    table.cell(1, 1).text = "必須"
    f = tmp_path / "flow.docx"
    doc.save(f)
    got = attachments.extract(f, "flow.docx")
    assert "業務フローの概要" in got.text
    assert "顧客名" in got.text and "必須" in got.text   # 表も拾う


def test_excelはシート名と値を読める(tmp_path: Path):
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "項目一覧"
    sheet.append(["列名", "型"])
    sheet.append(["顧客ID", "数値"])
    f = tmp_path / "items.xlsx"
    book.save(f)
    got = attachments.extract(f, "items.xlsx")
    assert "項目一覧" in got.text
    assert "顧客ID" in got.text


def test_excelはセル数の上限で打ち切る(tmp_path: Path, monkeypatch):
    from openpyxl import Workbook

    monkeypatch.setattr(attachments, "MAX_EXCEL_CELLS", 10)
    book = Workbook()
    sheet = book.active
    for i in range(50):
        sheet.append([f"値{i}", i])
    f = tmp_path / "big.xlsx"
    book.save(f)
    got = attachments.extract(f, "big.xlsx")
    assert "先頭" in got.note        # 打ち切ったことを利用者へ伝える
    assert "値49" not in got.text


def test_中身が空だとエラーになる(tmp_path: Path):
    from openpyxl import Workbook

    f = tmp_path / "empty.xlsx"
    Workbook().save(f)
    with pytest.raises(attachments.AttachmentError):
        attachments.extract(f, "empty.xlsx")


def test_文字が取れないpdfはスキャン扱いで断る(tmp_path: Path):
    """画像だけのPDFはOCRが要る。黙って空で通すと、資料を渡したのに
    設計へ反映されない、という分かりにくい失敗になる。"""
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    f = tmp_path / "scan.pdf"
    with open(f, "wb") as fp:
        writer.write(fp)
    with pytest.raises(attachments.AttachmentError) as exc:
        attachments.extract(f, "scan.pdf")
    assert "スキャン" in str(exc.value)


def test_長すぎるテキストは先頭だけ読む(tmp_path: Path, monkeypatch):
    monkeypatch.setattr(attachments, "MAX_EXTRACT_CHARS", 100)
    f = tmp_path / "long.txt"
    f.write_text("あ" * 500, encoding="utf-8")
    got = attachments.extract(f, "long.txt")
    assert len(got.text) == 100
    assert "先頭" in got.note
